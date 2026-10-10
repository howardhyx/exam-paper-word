import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

root = Path(__file__).parent
sys.path.insert(0, str(root.parent / 'skills' / 'exam-paper-word' / 'scripts'))
from word_components import add_answer_area, add_choices

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Mm(20)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12 if name == 'Normal' else 15)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), '宋体')
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + attr), None)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.25
    ppr = style.element.find(qn('w:pPr'))
    if ppr is not None:
        border = ppr.find(qn('w:pBdr'))
        if border is not None:
            ppr.remove(border)

plans = []


def space(label, **spec):
    paragraph = doc.add_paragraph(label)
    area = add_answer_area(doc, spec, question_paragraph=paragraph)
    plans.append({'label': label, 'configured': area['configured']})


doc.add_paragraph('答题区组件排版验证', style='Title')
doc.add_paragraph('验证连续空白、横线、作文格、绘图区、选项及双栏和跨页排版。')
space('1．数学推导 留连续空白60毫米', mode='blank', basis='5行推导与1行余量',
      estimated_lines=5, reserve_lines=1, line_height_mm=10)
space('2．文字回答 4条等长横线', mode='lined', basis='4行文字回答', estimated_lines=4)
p = doc.add_paragraph('3．四个选项分两行排列')
add_choices(doc, ['简短选项甲', '简短选项乙', '简短选项丙', '简短选项丁'], columns=2,
            question_paragraph=p)

doc.add_page_break()
doc.add_paragraph('作文格和绘图区验证', style='Title')
space('4．中文写作 105格，每格8毫米', mode='grid', basis='100字写作', estimated_chars=100,
      cell_size_mm=8)
space('5．几何作图 留连续空白50毫米', mode='drawing', basis='手画几何示意图', drawing_height_mm=50)
doc.add_paragraph('绘图区结束标记')

doc.add_page_break()
doc.add_paragraph('长作答区跨页验证', style='Title')
space('6．40行横线允许自然跨页', mode='lined', basis='多段文字作答', estimated_lines=40)
doc.add_paragraph('长作答区结束标记')

section = doc.add_section(WD_SECTION_START.NEW_PAGE)
section.page_width, section.page_height = Mm(420), Mm(297)
cols = section._sectPr.find(qn('w:cols'))
cols.set(qn('w:num'), '2')
cols.set(qn('w:space'), '567')
doc.add_paragraph('A3双栏验证', style='Title')
space('7．左栏横线按栏宽生成', mode='lined', basis='4行文字回答', estimated_lines=4)
p = doc.add_paragraph('8．较长选项使用每项一行')
add_choices(doc, ['较长的选项内容，允许自然换行，不把选项硬挤进很窄的单元格。'] * 4,
            columns=1, question_paragraph=p)
doc.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
space('9．右栏数学推导保持无横线', mode='blank', basis='5行推导', estimated_lines=5)
space('10．右栏文字回答使用横线', mode='lined', basis='4行文字回答', estimated_lines=4)

generated = root / 'generated'
generated.mkdir(exist_ok=True)
output = generated / 'component-demo.docx'
doc.core_properties.author = 'howardhyx'
doc.core_properties.last_modified_by = ''
doc.save(output)
(generated / 'component-demo-configured.json').write_text(json.dumps(plans, ensure_ascii=False, indent=2), encoding='utf-8')
print(output)
