#!/usr/bin/env python3
"""Reusable python-docx exam spaces and choice layouts; no full-paper template."""

import math

from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from answer_space import numeric, plan_space


def _set(parent, tag, **attributes):
    child = parent.find(qn("w:" + tag))
    if child is None:
        child = OxmlElement("w:" + tag)
        parent.append(child)
    for name, value in attributes.items():
        child.set(qn("w:" + name), str(value))
    return child


def _twips(mm):
    return round(mm * 1440 / 25.4)


def _font(rpr, east_asia, latin, size_pt):
    fonts = _set(rpr, "rFonts", ascii=latin, hAnsi=latin, eastAsia=east_asia, cs=latin)
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
        fonts.attrib.pop(qn("w:" + name), None)
    _set(rpr, "lang", val="en-US", eastAsia="zh-CN")
    _set(rpr, "sz", val=round(size_pt * 2))
    _set(rpr, "szCs", val=round(size_pt * 2))
    _set(rpr, "color", val="000000")


def _paragraph(paragraph, east_asia, latin, size_pt, *, empty=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.left_indent = Mm(0)
    fmt.right_indent = Mm(0)
    fmt.first_line_indent = Mm(0)
    fmt.keep_with_next = False
    fmt.keep_together = False
    fmt.widow_control = False
    fmt.line_spacing = Pt(size_pt) if empty else 1.25
    ppr = paragraph._p.get_or_add_pPr()
    edges = _set(ppr, "pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        _set(edges, edge, val="nil")
    _font(_set(ppr, "rPr"), east_asia, latin, size_pt)
    for run in paragraph.runs:
        _font(run._r.get_or_add_rPr(), east_asia, latin, size_pt)


def available_width_mm(section, *, column_index=0):
    """Use the current section and column; custom unequal columns are supported."""
    full = (section.page_width - section.left_margin - section.right_margin - section.gutter) / 36000
    return _column_width(section, full, column_index)


def _column_width(section, full, column_index):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        if column_index != 0:
            raise ValueError("single-column section has only column_index=0")
        return full
    explicit = cols.findall(qn("w:col"))
    count = int(cols.get(qn("w:num"), len(explicit) or 1))
    if not 0 <= column_index < count:
        raise ValueError("column_index is out of range")
    if cols.get(qn("w:equalWidth"), "1") in {"0", "false"}:
        if len(explicit) != count:
            raise ValueError("unequal column widths must be explicit")
        return int(explicit[column_index].get(qn("w:w"))) * 25.4 / 1440
    gap = int(cols.get(qn("w:space"), "720")) * 25.4 / 1440
    return (full - (count - 1) * gap) / count


def _width(doc, section, width_mm, left_mm, right_mm, column_index):
    section = section or doc.sections[-1]
    full = (section.page_width - section.left_margin - section.right_margin - section.gutter) / 36000
    available = _column_width(section, full, column_index)
    left = numeric(left_mm, "left_mm")
    right = numeric(right_mm, "right_mm")
    maximum = available - left - right
    width = maximum if width_mm is None else numeric(width_mm, "width_mm", positive=True)
    if maximum <= 0 or width > maximum + 0.01:
        raise ValueError("component width exceeds the selected column after indentation")
    return section, width, left


def _table(doc, rows, columns, width_mm, left_mm, *, borders="none", padding_mm=0):
    table = doc.add_table(rows=rows, cols=columns)
    style_name = "ExamPaperWordPlainComponent"
    if style_name not in doc.styles:
        neutral_style = doc.styles.add_style(style_name, WD_STYLE_TYPE.TABLE)
        neutral_style.base_style = None
    table.style = doc.styles[style_name]
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    pr = table._tbl.tblPr
    # Use an isolated style instead of inheriting the document's default table style.
    _set(pr, "tblW", w=_twips(width_mm), type="dxa")
    _set(pr, "tblInd", w=_twips(left_mm), type="dxa")
    _set(pr, "tblLayout", type="fixed")
    # Absence preserves collapsed shared borders; explicit zero spacing can
    # still produce separate cell-edge strokes in some native Word exporters.
    spacing = pr.find(qn("w:tblCellSpacing"))
    if spacing is not None:
        pr.remove(spacing)
    _set(pr, "tblLook", val="0000", firstRow=0, lastRow=0, firstColumn=0,
         lastColumn=0, noHBand=1, noVBand=1)
    edges = _set(pr, "tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _set(edges, edge, val="single" if borders == "grid" else "nil", sz=4, color="000000")
    margins = _set(pr, "tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        _set(margins, side, w=_twips(padding_mm), type="dxa")
    # Cell margins, table grid and preferred widths must agree.
    for col in table.columns:
        col.width = Mm(width_mm / columns)
    for row in table.rows:
        _set(row._tr.get_or_add_trPr(), "cantSplit")
        for cell in row.cells:
            cell.width = Mm(width_mm / columns)
            cellpr = cell._tc.get_or_add_tcPr()
            _set(cellpr, "tcW", w=_twips(width_mm / columns), type="dxa")
            cm = _set(cellpr, "tcMar")
            for side in ("top", "left", "bottom", "right"):
                _set(cm, side, w=_twips(padding_mm), type="dxa")
            # Direct cell borders override even a document's default table style.
            if borders != "grid":
                cb = _set(cellpr, "tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    visible = borders == "lined" and edge == "bottom"
                    _set(cb, edge, val="single" if visible else "nil", sz=4, color="000000")
            _set(cellpr, "shd", val="clear", fill="auto")
    return table


def add_answer_area(doc, spec, *, section=None, width_mm=None, left_mm=0,
                    right_mm=0, column_index=0, east_asia="宋体", latin="Times New Roman",
                    font_size_pt=12, question_paragraph=None):
    """Append a native answer area, returning tables and configured capacity.

    blank = invisible row structure; lined = horizontal rules only;
    grid = square writing cells; drawing = uninterrupted blank drawing space.
    Configured capacity is never a substitute for rendered page observations.
    """
    if not isinstance(spec, dict):
        raise ValueError("answer area specification must be an object")
    if spec.get("mode") == "none":
        return {"tables": [], "configured": plan_space(spec, None)}
    requested_width = spec.get("width_mm") if width_mm is None else width_mm
    section, width, left = _width(doc, section, requested_width, left_mm, right_mm, column_index)
    plan = plan_space(spec, width)
    height_available = (section.page_height - section.top_margin - section.bottom_margin) / 36000
    if plan.get("drawing_height_mm", 0) > height_available:
        raise ValueError("drawing area cannot fit one page; divide the task or use another page size")
    if plan.get("row_height_mm", 0) > height_available:
        raise ValueError("one handwriting row cannot fit the selected page")
    if question_paragraph is not None:
        question_paragraph.paragraph_format.keep_with_next = True
    tables = []
    if plan["mode"] in {"blank", "lined", "grid"}:
        grid = plan["mode"] == "grid"
        table = _table(doc, plan["writing_lines"], plan["columns"] if grid else 1,
                       plan["width_mm"], left,
                       borders="grid" if grid else ("lined" if plan["mode"] == "lined" else "none"))
        for row in table.rows:
            row.height = Mm(plan["row_height_mm"])
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                _paragraph(cell.paragraphs[0], east_asia, latin, font_size_pt, empty=True)
        tables.append(table)
    if plan.get("drawing_height_mm", 0):
        table = _table(doc, 1, 1, width, left)
        table.rows[0].height = Mm(plan["drawing_height_mm"])
        table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        _paragraph(table.cell(0, 0).paragraphs[0], east_asia, latin, font_size_pt, empty=True)
        tables.append(table)
    return {"tables": tables, "configured": plan}


def add_choices(doc, choices, *, columns=2, section=None, width_mm=None,
                left_mm=0, right_mm=0, column_index=0, east_asia="宋体",
                latin="Times New Roman", font_size_pt=12, question_paragraph=None):
    """Use 1, 2 or 4 columns selected after inspecting option width.

    Strings receive A．/B． labels. A callable can populate a cell with native
    formulas/images; it must supply its own label. No LaTeX text conversion here.
    """
    if isinstance(columns, bool) or columns not in {1, 2, 4} or not choices or len(choices) > 26:
        raise ValueError("use nonempty choices (up to 26) and 1, 2 or 4 columns")
    section, width, left = _width(doc, section, width_mm, left_mm, right_mm, column_index)
    table = _table(doc, math.ceil(len(choices) / columns), columns, width, left, padding_mm=1)
    if question_paragraph is not None:
        question_paragraph.paragraph_format.keep_with_next = True
    for index, cell in enumerate(c for row in table.rows for c in row.cells):
        _paragraph(cell.paragraphs[0], east_asia, latin, font_size_pt)
        if index < len(choices):
            item = choices[index]
            if callable(item):
                item(cell)
            elif isinstance(item, str):
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.left_indent = Mm(6)
                paragraph.paragraph_format.first_line_indent = Mm(-6)
                run = paragraph.add_run(f"{chr(65 + index)}．{item}")
                _font(run._r.get_or_add_rPr(), east_asia, latin, font_size_pt)
            else:
                raise ValueError("choice must be text or a cell-populating callable")
    return table
