import copy
import sys
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'skills' / 'exam-paper-word' / 'scripts'))
from answer_space import plan_space, audit_spaces
from word_components import add_answer_area, add_choices, available_width_mm


def document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(20)
    return doc


class ComponentTests(unittest.TestCase):
    def test_minimum_space_and_deficit(self):
        spec = dict(mode='blank', basis='five handwritten steps', estimated_lines=5,
                    line_height_mm=10, reserve_lines=1)
        self.assertEqual(plan_space(spec, 170)['height_mm'], 60)
        with self.assertRaisesRegex(ValueError, 'below required'):
            plan_space(dict(spec, allocated_height_mm=40), 170)

    def test_blank_has_no_visible_borders_or_underscores(self):
        doc = document()
        result = add_answer_area(doc, dict(mode='blank', basis='推导', estimated_lines=6))
        table = result['tables'][0]
        self.assertAlmostEqual(result['configured']['width_mm'], 170, places=2)
        self.assertEqual(len(table.rows), 6)
        self.assertTrue(all(edge.get(qn('w:val')) == 'nil'
                            for edge in table._tbl.tblPr.find(qn('w:tblBorders'))))
        self.assertEqual(''.join(table._tbl.itertext()).strip(), '')
        self.assertNotIn('_', table.cell(0, 0).text)

    def test_full_width_lined_area_with_indentation(self):
        doc = document()
        result = add_answer_area(doc, dict(mode='lined', basis='文字回答', estimated_lines=4), left_mm=8)
        table = result['tables'][0]
        self.assertAlmostEqual(result['configured']['width_mm'], 162, delta=.02)
        self.assertEqual([c.width for row in table.rows for c in row.cells],
                         [table.cell(0, 0).width] * 4)
        borders = table._tbl.tblPr.find(qn('w:tblBorders'))
        self.assertEqual(borders.find(qn('w:left')).get(qn('w:val')), 'nil')
        cell_borders = table.cell(0, 0)._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
        self.assertEqual(cell_borders.find(qn('w:bottom')).get(qn('w:val')), 'single')
        self.assertEqual(cell_borders.find(qn('w:top')).get(qn('w:val')), 'nil')

    def test_grid_capacity_and_square_cells(self):
        doc = document()
        result = add_answer_area(doc, dict(mode='grid', basis='中文写作', estimated_chars=100, cell_size_mm=8))
        self.assertEqual(result['configured']['grid_cells'], 105)
        self.assertEqual(result['configured']['width_mm'], 168)
        self.assertAlmostEqual(result['tables'][0].cell(0, 0).width.mm, 8, places=1)
        self.assertAlmostEqual(result['tables'][0].rows[0].height.mm, 8, places=1)

    def test_a3_columns_and_overwide_component(self):
        doc = document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Mm(420), Mm(297)
        cols = sec._sectPr.find(qn('w:cols'))
        cols.set(qn('w:num'), '2')
        cols.set(qn('w:space'), '567')  # approximately 10 mm
        self.assertAlmostEqual(available_width_mm(sec), 185, delta=.02)
        with self.assertRaisesRegex(ValueError, 'exceeds'):
            add_answer_area(doc, dict(mode='blank', basis='推导', estimated_lines=3), width_mm=380)
        cols.set(qn('w:equalWidth'), '0')
        for width in ('9000', '11976'):
            child = OxmlElement('w:col')
            child.set(qn('w:w'), width)
            cols.append(child)
        self.assertAlmostEqual(available_width_mm(sec, column_index=0), 158.75)

    def test_drawing_and_choices(self):
        doc = document()
        area = add_answer_area(doc, dict(mode='drawing', basis='几何作图', drawing_height_mm=50))
        self.assertEqual(len(area['tables'][0].rows), 1)
        self.assertAlmostEqual(area['tables'][0].rows[0].height.mm, 50, places=1)
        table = add_choices(doc, ['甲', '乙', '丙', '丁'], columns=2)
        self.assertEqual([[cell.text for cell in row.cells] for row in table.rows],
                         [['A．甲', 'B．乙'], ['C．丙', 'D．丁']])
        with self.assertRaisesRegex(ValueError, 'cannot fit'):
            add_answer_area(doc, dict(mode='drawing', basis='超高图', drawing_height_mm=300))

    def test_final_observation_required_and_deficit_rejected(self):
        space = dict(id='1(1)', mode='blank', basis='推导', width_mm=170,
                     estimated_lines=5, line_height_mm=10, reserve_lines=1)
        data = {'questions': [{'id': '1', 'answer_spaces': [space]}, {'id': '2', 'answer_spaces': []}]}
        self.assertTrue(audit_spaces(data, 'plan')['ok'])
        self.assertFalse(audit_spaces(data, 'final')['ok'])
        measured = copy.deepcopy(data)
        measured['questions'][0]['answer_spaces'][0]['observed'] = dict(height_mm=60, width_mm=170, writing_lines=6)
        self.assertTrue(audit_spaces(measured, 'final')['ok'])
        measured['questions'][0]['answer_spaces'][0]['observed']['writing_lines'] = 4
        self.assertFalse(audit_spaces(measured, 'final')['ok'])


if __name__ == '__main__':
    unittest.main(verbosity=2)
